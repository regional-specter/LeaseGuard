"""Tests for source manifests, PDF/DOCX parsing, and the document pipeline."""

import hashlib
from datetime import date
from pathlib import Path
from typing import Any
from unittest.mock import patch

import pytest
from docx import Document
from pydantic import ValidationError

from leaseguard.evaluation.paths import repo_root
from leaseguard.ingestion.acquire import (
    DocumentAcquisitionError,
    DocumentDownloadError,
    acquire_source,
    raw_destination,
)
from leaseguard.ingestion.cli import main
from leaseguard.ingestion.manifest import get_source, load_pipeline_config, load_source_manifest
from leaseguard.ingestion.models import PipelineConfig, SourceEntry, SourceManifest
from leaseguard.ingestion.parse import parse_document
from leaseguard.ingestion.pipeline import process_manifest
from leaseguard.ingestion.validate import (
    DocumentValidationError,
    detect_media_type,
    sha256_file,
    validate_acquired_file,
)
from leaseguard.ontology.enums import DocumentType, RightsStatus
from leaseguard.preprocessing.duplicates import (
    exact_duplicate_map,
    jaccard_similarity,
    near_duplicate_map,
    word_shingles,
)
from leaseguard.preprocessing.models import PageRecord, ProcessedDocument, TextBlock
from leaseguard.preprocessing.normalize import assign_offsets, join_pages, normalize_whitespace
from leaseguard.preprocessing.segment import segment_pages
from leaseguard.privacy import redact_personal_identifiers

ROOT = repo_root()


def build_simple_pdf(text: str) -> bytes:
    """Build a one-page PDF that pypdf can extract text from."""
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1", "replace")
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
        (
            b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n"
        ),
        b"4 0 obj << /Length %d >> stream\n" % len(stream) + stream + b"\nendstream endobj\n",
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
    ]
    header = b"%PDF-1.4\n"
    body = header
    offsets: list[int] = []
    for obj in objects:
        offsets.append(len(body))
        body += obj
    xref_start = len(body)
    xref = [b"xref\n0 6\n0000000000 65535 f \n"]
    xref.extend(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    trailer = (
        b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref_start).encode() + b"\n%%EOF\n"
    )
    return body + b"".join(xref) + trailer


def write_docx(path: Path) -> None:
    """Write a tiny office-lease DOCX with a heading and table."""
    document = Document()
    document.add_heading("Article 1 Rent", level=1)
    document.add_paragraph("Tenant shall pay Landlord base rent each month.")
    document.add_heading("Exhibit A Premises", level=1)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Suite"
    table.rows[0].cells[1].text = "1200"
    document.save(str(path))


def source_entry(tmp_path: Path, **changes: Any) -> tuple[SourceEntry, Path]:
    """Create a markdown source file and matching manifest entry."""
    text = changes.pop("text", "Section 1. Rent\nTenant shall pay rent of 5000 USD.")
    requested = changes.get("local_path")
    path = Path(requested) if requested else tmp_path / "lease.md"
    if not path.is_absolute():
        path = tmp_path / path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    digest = sha256_file(path)
    values: dict[str, Any] = {
        "document_id": "sample-lease",
        "family_id": "sample-family",
        "document_type": DocumentType.OFFICE_LEASE,
        "source_name": "Synthetic sample",
        "source_url": "https://example.invalid/lease.md",
        "local_path": str(path),
        "sha256": digest,
        "rights_status": RightsStatus.OPEN_LICENSE,
        "source_license": "CC0-1.0",
        "media_type": "markdown",
        "training_eligible": False,
        "retrieved_at": date(2026, 9, 8),
    }
    values.update(changes)
    values["local_path"] = str(path)
    return SourceEntry(**values), path


def pipeline_config() -> PipelineConfig:
    """Return local test pipeline limits."""
    return PipelineConfig(
        pipeline_version="1.0.0",
        allowed_media_types=["pdf", "docx", "markdown", "text"],
        max_file_bytes=1_000_000,
        near_duplicate_threshold=0.5,
        ocr_enabled=False,
        allow_download_env="LEASEGUARD_ALLOW_DOCUMENT_DOWNLOAD",
        shingle_size=3,
    )


def test_phase2_manifest_is_a_valid_source_manifest() -> None:
    """The frozen validation sources are the Phase 4 approved source list."""
    manifest = load_source_manifest(ROOT / "data" / "manifests" / "phase2_validation_sources.json")
    assert len(manifest.sources) == 4
    gsa = get_source(manifest, "gsa-l100-2026")
    assert load_pipeline_config(ROOT / "configs" / "data" / "pipeline.v1.json").ocr_enabled is False
    assert gsa.media_type == "pdf"
    assert "gsa-amendment-template-2021" in gsa.related_document_ids
    assert gsa.training_eligible is False
    with pytest.raises(KeyError, match="Unknown document_id"):
        get_source(manifest, "missing")


def test_source_entry_rejects_insecure_urls_and_missing_rights() -> None:
    """HTTPS and rights metadata are required before a source can be acquired."""
    with pytest.raises(ValidationError, match="https"):
        SourceEntry(
            document_id="x",
            family_id="y",
            document_type=DocumentType.OFFICE_LEASE,
            source_name="Example",
            source_url="http://example.invalid/file.pdf",
            local_path="file.pdf",
            sha256="a" * 64,
            rights_status=RightsStatus.PUBLIC_DOMAIN,
            media_type="pdf",
        )
    with pytest.raises(ValidationError, match="source_license"):
        SourceEntry(
            document_id="x",
            family_id="y",
            document_type=DocumentType.OFFICE_LEASE,
            source_name="Example",
            local_path="file.md",
            sha256="a" * 64,
            rights_status=RightsStatus.OPEN_LICENSE,
            media_type="markdown",
        )
    with pytest.raises(ValidationError, match="rights_notes"):
        SourceEntry(
            document_id="x",
            family_id="y",
            document_type=DocumentType.OFFICE_LEASE,
            source_name="Example",
            local_path="file.md",
            sha256="a" * 64,
            rights_status=RightsStatus.REVIEW_REQUIRED,
            media_type="markdown",
        )


def test_detects_pdf_docx_and_markdown(tmp_path: Path) -> None:
    """File identity comes from magic bytes, not only the filename."""
    pdf_path = tmp_path / "a.pdf"
    pdf_path.write_bytes(build_simple_pdf("Tenant shall pay rent."))
    docx_path = tmp_path / "a.docx"
    write_docx(docx_path)
    md_path = tmp_path / "a.md"
    md_path.write_text("# Lease\n", encoding="utf-8")
    txt_path = tmp_path / "a.txt"
    txt_path.write_text("plain", encoding="utf-8")
    bin_path = tmp_path / "a.bin"
    bin_path.write_bytes(b"not-a-document")

    assert detect_media_type(pdf_path) == "pdf"
    assert detect_media_type(docx_path) == "docx"
    assert detect_media_type(md_path) == "markdown"
    assert detect_media_type(txt_path) == "text"
    with pytest.raises(DocumentValidationError, match="unsupported"):
        detect_media_type(bin_path)


def test_checksum_and_size_validation(tmp_path: Path) -> None:
    """A file that does not match the manifest is rejected before parsing."""
    source, path = source_entry(tmp_path)
    config = pipeline_config()
    assert validate_acquired_file(path, source, config) == "markdown"
    wrong = source.model_copy(update={"sha256": "b" * 64})
    with pytest.raises(DocumentValidationError, match="checksum"):
        validate_acquired_file(path, wrong, config)
    tiny = config.model_copy(update={"max_file_bytes": 1})
    with pytest.raises(DocumentValidationError, match="exceeds"):
        validate_acquired_file(path, source, tiny)


def test_local_import_does_not_overwrite_verified_raw(tmp_path: Path) -> None:
    """A verified raw copy is reused; a conflicting copy is refused."""
    source, _path = source_entry(tmp_path)
    raw_root = tmp_path / "raw"
    first = acquire_source(
        source,
        raw_root=raw_root,
        search_roots=[tmp_path],
        config=pipeline_config(),
        allow_download=False,
        download_env=None,
    )
    assert first.is_file()
    again = acquire_source(
        source,
        raw_root=raw_root,
        search_roots=[tmp_path],
        config=pipeline_config(),
        allow_download=False,
        download_env=None,
    )
    assert again == first
    dest = raw_destination(raw_root, source)
    dest.write_text("tampered", encoding="utf-8")
    with pytest.raises(DocumentAcquisitionError, match="does not match"):
        acquire_source(
            source,
            raw_root=raw_root,
            search_roots=[tmp_path],
            config=pipeline_config(),
            allow_download=False,
            download_env=None,
        )


def test_downloads_are_blocked_locally_and_allowed_when_gated(tmp_path: Path) -> None:
    """Colab must set the download gate; the local machine must not fetch sources."""
    pdf_bytes = build_simple_pdf("Base rent is due monthly.")
    source = SourceEntry(
        document_id="remote-lease",
        family_id="remote-family",
        document_type=DocumentType.OFFICE_LEASE,
        source_name="Remote",
        source_url="https://example.invalid/lease.pdf",
        local_path="missing.pdf",
        sha256=hashlib.sha256(pdf_bytes).hexdigest(),
        rights_status=RightsStatus.PUBLIC_DOMAIN,
        media_type="pdf",
    )
    with pytest.raises(DocumentDownloadError, match="blocked"):
        acquire_source(
            source,
            raw_root=tmp_path / "raw",
            search_roots=[tmp_path],
            config=pipeline_config(),
            allow_download=False,
            download_env=None,
        )
    orphan = source.model_copy(update={"source_url": None, "local_path": "nowhere.pdf"})
    with pytest.raises(DocumentAcquisitionError, match="no local file"):
        acquire_source(
            orphan,
            raw_root=tmp_path / "raw",
            search_roots=[tmp_path],
            config=pipeline_config(),
            allow_download=True,
            download_env="1",
        )

    class _Response:
        def __enter__(self) -> object:
            return self

        def __exit__(self, *args: object) -> None:
            return None

        def read(self) -> bytes:
            return pdf_bytes

    with patch("leaseguard.ingestion.acquire.urlopen", return_value=_Response()):
        downloaded = acquire_source(
            source,
            raw_root=tmp_path / "raw",
            search_roots=[tmp_path],
            config=pipeline_config(),
            allow_download=True,
            download_env="1",
        )
    assert downloaded.is_file()
    assert sha256_file(downloaded) == source.sha256


def test_parses_pdf_docx_and_segments_offsets(tmp_path: Path) -> None:
    """Digital parsers preserve page or heading structure and exact offsets."""
    pdf_path = tmp_path / "lease.pdf"
    pdf_path.write_bytes(build_simple_pdf("Tenant shall pay rent."))
    pdf_pages = parse_document(pdf_path, "pdf")
    assert pdf_pages[0].page_number == 1
    assert "rent" in pdf_pages[0].text.lower()

    docx_path = tmp_path / "lease.docx"
    write_docx(docx_path)
    docx_pages = parse_document(docx_path, "docx")
    normalized, blocks = segment_pages(docx_pages)
    assign_offsets(blocks, normalized)
    kinds = {block.kind for block in blocks}
    assert "heading" in kinds
    assert "exhibit" in kinds
    assert "table" in kinds
    assert normalized[blocks[0].start_char : blocks[0].end_char] == blocks[0].text


def test_redaction_and_duplicate_helpers() -> None:
    """Processed text can be redacted and compared without touching raw files."""
    result = redact_personal_identifiers(
        "Email jane@example.com phone 415-555-0100 ssn 123-45-6789."
    )
    assert "[REDACTED_EMAIL]" in result.text
    assert "[REDACTED_PHONE]" in result.text
    assert "[REDACTED_SSN]" in result.text
    assert set(result.codes) == {"email", "phone", "ssn"}
    assert word_shingles("one two three", 5) == {"one two three"}
    assert jaccard_similarity(set(), set()) == 1.0
    assert jaccard_similarity({"a"}, set()) == 0.0
    joined = join_pages(
        [
            PageRecord(page_number=1, text="First page"),
            PageRecord(page_number=2, text="Second page"),
        ]
    )
    assert "First page" in joined and "Second page" in joined
    normalized, blocks = segment_pages(
        [
            PageRecord(page_number=1, text="   "),
            PageRecord(page_number=2, text="Section 4. Notices\nTenant shall send notice."),
        ]
    )
    assert blocks
    assert "Notices" in normalized


def test_process_manifest_exports_records_and_resumes(tmp_path: Path) -> None:
    """An approved source becomes a validated intermediate record and can resume."""
    first, path = source_entry(tmp_path, document_id="lease-a")
    near_text = "Section 1. Rent\nTenant shall pay rent of 5000 USD each month promptly."
    second, _ = source_entry(
        tmp_path,
        document_id="lease-b",
        family_id="sample-family",
        text=near_text,
        local_path=str(tmp_path / "lease-b.md"),
    )
    third, _ = source_entry(
        tmp_path,
        document_id="lease-c",
        family_id="other-family",
        local_path=str(tmp_path / "lease-c.md"),
        text=path.read_text(encoding="utf-8"),
    )
    manifest = SourceManifest(
        manifest_version="1.0.0",
        purpose="pipeline tests",
        storage_policy="temporary test files",
        sources=[
            first.model_copy(update={"related_document_ids": ["lease-b"]}),
            second,
            third,
        ],
    )
    raw_root = tmp_path / "raw"
    processed_root = tmp_path / "processed"
    checkpoint_dir = tmp_path / "checkpoints"
    batch = process_manifest(
        manifest,
        config=pipeline_config(),
        raw_root=raw_root,
        processed_root=processed_root,
        checkpoint_dir=checkpoint_dir,
        search_roots=[tmp_path],
        allow_download=False,
        download_env=None,
    )
    assert {item.status for item in batch.reports} == {"exported"}
    exported = ProcessedDocument.model_validate_json(
        (processed_root / "lease-a.json").read_text(encoding="utf-8")
    )
    assert exported.source.document_id == "lease-a"
    assert exported.blocks
    raw_file = Path(exported.raw_path)
    original = sha256_file(raw_file)
    resumed = process_manifest(
        manifest,
        config=pipeline_config(),
        raw_root=raw_root,
        processed_root=processed_root,
        checkpoint_dir=checkpoint_dir,
        search_roots=[tmp_path],
        allow_download=False,
        download_env=None,
    )
    assert {item.status for item in resumed.reports} == {"skipped"}
    assert sha256_file(raw_file) == original
    by_id = {item.document_id: item for item in batch.reports}
    assert by_id["lease-c"].duplicate_of == "lease-a"
    assert "lease-a" in by_id["lease-b"].near_duplicate_of
    family_ids = {group.family_id for group in batch.families}
    assert "sample-family" in family_ids


def test_process_reports_failures_and_cli_roundtrip(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Missing files fail the document, and the CLI validates the frozen manifest."""
    assert main(["validate-manifest"]) == 0
    assert "4 sources" in capsys.readouterr().out

    source, _path = source_entry(tmp_path, document_id="broken")
    missing = source.model_copy(update={"local_path": str(tmp_path / "absent.md")})
    manifest = SourceManifest(
        manifest_version="1.0.0",
        purpose="failure",
        storage_policy="test",
        sources=[missing],
    )
    manifest_path = tmp_path / "manifest.json"
    manifest_path.write_text(manifest.model_dump_json(), encoding="utf-8")
    config_path = ROOT / "configs" / "data" / "pipeline.v1.json"
    assert (
        main(
            [
                "process",
                "--manifest",
                str(manifest_path),
                "--config",
                str(config_path),
                "--raw-root",
                str(tmp_path / "raw"),
                "--processed-root",
                str(tmp_path / "processed"),
                "--checkpoint-dir",
                str(tmp_path / "checkpoints"),
                "--search-root",
                str(tmp_path),
                "--output",
                str(tmp_path / "report.json"),
            ]
        )
        == 1
    )
    assert "failed" in capsys.readouterr().err

    good_manifest = SourceManifest(
        manifest_version="1.0.0",
        purpose="success",
        storage_policy="test",
        sources=[source],
    )
    good_path = tmp_path / "good.json"
    good_path.write_text(good_manifest.model_dump_json(), encoding="utf-8")
    assert (
        main(
            [
                "process",
                "--manifest",
                str(good_path),
                "--raw-root",
                str(tmp_path / "raw2"),
                "--processed-root",
                str(tmp_path / "processed2"),
                "--checkpoint-dir",
                str(tmp_path / "checkpoints2"),
                "--search-root",
                str(tmp_path),
            ]
        )
        == 0
    )


def test_duplicate_manifest_ids_and_normalize() -> None:
    """Manifest identifiers stay unique and whitespace normalization is stable."""
    payload = load_source_manifest(
        ROOT / "data" / "manifests" / "phase2_validation_sources.json"
    ).model_dump()
    payload["sources"].append(payload["sources"][0])
    with pytest.raises(ValidationError, match="unique"):
        SourceManifest.model_validate(payload)
    assert normalize_whitespace("  Tenant   shall\n  pay  ") == "Tenant shall\npay"


def test_exact_and_near_duplicate_maps_on_processed_records(tmp_path: Path) -> None:
    """Duplicate helpers operate on processed records rather than raw paths."""
    source, _ = source_entry(tmp_path)
    batch = process_manifest(
        SourceManifest(
            manifest_version="1.0.0",
            purpose="dup",
            storage_policy="test",
            sources=[source],
        ),
        config=pipeline_config(),
        raw_root=tmp_path / "raw",
        processed_root=tmp_path / "processed",
        checkpoint_dir=tmp_path / "checkpoints",
        search_roots=[tmp_path],
        allow_download=False,
        download_env=None,
    )
    document = ProcessedDocument.model_validate_json(
        Path(batch.reports[0].output_path or "").read_text(encoding="utf-8")
    )
    clone = document.model_copy(
        update={"source": document.source.model_copy(update={"document_id": "clone"})}
    )
    assert exact_duplicate_map([document, clone])["clone"] == "sample-lease"
    near = near_duplicate_map(
        [document, clone],
        threshold=0.5,
        shingle_size=3,
    )
    assert "sample-lease" in near["clone"]


def test_empty_file_media_mismatch_and_bad_offsets(tmp_path: Path) -> None:
    """Parser guards reject empty files, type mismatches, and broken offsets."""
    empty = tmp_path / "empty.md"
    empty.write_bytes(b"")
    source, _ = source_entry(tmp_path)
    empty_source = source.model_copy(
        update={"local_path": str(empty), "sha256": sha256_file(empty)}
    )
    with pytest.raises(DocumentValidationError, match="empty"):
        validate_acquired_file(empty, empty_source, pipeline_config())
    pdf_declared = source.model_copy(update={"media_type": "pdf"})
    with pytest.raises(DocumentValidationError, match="detected markdown"):
        validate_acquired_file(Path(source.local_path), pdf_declared, pipeline_config())
    with pytest.raises(ValueError, match="offsets"):
        assign_offsets(
            [
                TextBlock(
                    block_id="block-0001",
                    kind="paragraph",
                    text="nope",
                    start_char=0,
                    end_char=4,
                )
            ],
            "abcd",
        )


def test_download_rejects_checksum_mismatch(tmp_path: Path) -> None:
    """A downloaded payload that does not match the manifest is discarded."""
    expected = build_simple_pdf("expected text")
    source = SourceEntry(
        document_id="remote-bad",
        family_id="remote-family",
        document_type=DocumentType.OFFICE_LEASE,
        source_name="Remote",
        source_url="https://example.invalid/lease.pdf",
        local_path="missing.pdf",
        sha256=hashlib.sha256(expected).hexdigest(),
        rights_status=RightsStatus.PUBLIC_DOMAIN,
        media_type="pdf",
    )

    class _Response:
        def __enter__(self) -> object:
            return self

        def __exit__(self, *args: object) -> None:
            return None

        def read(self) -> bytes:
            return build_simple_pdf("other text")

    with (
        patch("leaseguard.ingestion.acquire.urlopen", return_value=_Response()),
        pytest.raises(DocumentAcquisitionError, match="does not match"),
    ):
        acquire_source(
            source,
            raw_root=tmp_path / "raw",
            search_roots=[tmp_path],
            config=pipeline_config(),
            allow_download=True,
            download_env="1",
        )
