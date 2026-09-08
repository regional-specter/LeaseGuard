"""Read digital PDF, DOCX, markdown, and plain-text files."""

from collections.abc import Iterator
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from leaseguard.ingestion.models import MediaType
from leaseguard.preprocessing.models import PageRecord

HEADING_STYLES = {"Heading 1", "Heading 2", "Heading 3", "Title"}


class DocumentParseError(ValueError):
    """Raised when a digital file cannot be read as text."""


def parse_document(path: Path, media_type: MediaType) -> list[PageRecord]:
    """Extract page-ordered text without modifying the raw file."""
    if media_type == "pdf":
        return _parse_pdf(path)
    if media_type == "docx":
        return _parse_docx(path)
    if media_type in {"markdown", "text"}:
        return _parse_text(path)
    raise DocumentParseError(f"unsupported media type: {media_type}")


def _parse_pdf(path: Path) -> list[PageRecord]:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise DocumentParseError(f"{path.name} is encrypted")
    pages: list[PageRecord] = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        warnings: list[str] = []
        if not text.strip():
            warnings.append("no extractable text; OCR is deferred")
        pages.append(PageRecord(page_number=index, text=text, warnings=warnings))
    if not pages:
        raise DocumentParseError(f"{path.name} contains no pages")
    return pages


def _parse_docx(path: Path) -> list[PageRecord]:
    document = docx.Document(str(path))
    lines: list[str] = []
    for item in _iter_docx_items(document):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            if not text:
                continue
            if item.style is not None and item.style.name in HEADING_STYLES:
                lines.append(f"# {text}")
            else:
                lines.append(text)
        elif isinstance(item, Table):
            rendered = _render_table(item)
            if rendered:
                lines.append(rendered)
    body = "\n".join(lines)
    warnings: list[str] = []
    if not body.strip():
        warnings.append("no extractable text; OCR is deferred")
    return [PageRecord(page_number=1, text=body, warnings=warnings)]


def _iter_docx_items(document: DocxDocument) -> Iterator[Paragraph | Table]:
    iterate = getattr(document, "iter_inner_content", None)
    if iterate is not None:
        yield from iterate()
        return
    yield from document.paragraphs
    yield from document.tables


def _parse_text(path: Path) -> list[PageRecord]:
    text = path.read_text(encoding="utf-8")
    return [PageRecord(page_number=1, text=text, warnings=[])]


def _render_table(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)
